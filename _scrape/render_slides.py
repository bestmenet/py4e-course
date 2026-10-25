#!/usr/bin/env python3
"""Turn every lecture deck into web-native content.

For each Pythonlearn-*.pdf we render one PNG per page (so the deck can be
paged through inside the site with no PDF plugin) and for each matching
.pptx we pull the structured text (title / bullets / speaker notes) which
is cleaner than raw PDF text extraction.  Result -> ../data/slides.json
"""
import glob
import json
import os
import re
import sys
import time

import fitz  # PyMuPDF
from pptx import Presentation

ROOT = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(ROOT, "..", "data")
STATIC = os.path.join(ROOT, "..", "static")
DECKS = os.path.join(STATIC, "slides", "downloads")

ZOOM = 2.0  # 720x540pt slides -> 1440x1080px


def deck_slug(name):
    m = re.match(r"Pythonlearn-(\d+)-(.+)", name)
    if m:
        return "%s-%s" % (m.group(1), re.sub(r"[^A-Za-z0-9]+", "-", m.group(2)).strip("-").lower())
    return re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").lower()


def render_pdf(path, slug):
    out_dir = os.path.join(STATIC, "slides", slug)
    os.makedirs(out_dir, exist_ok=True)
    doc = fitz.open(path)
    pages = []
    for i, page in enumerate(doc, start=1):
        png = os.path.join(out_dir, "slide-%02d.png" % i)
        if not (os.path.exists(png) and os.path.getsize(png) > 0):
            pix = page.get_pixmap(matrix=fitz.Matrix(ZOOM, ZOOM), alpha=False)
            pix.save(png)
        text = page.get_text("text")
        # PDF text is useful as a fallback / search source: collapse whitespace
        text = "\n".join(line.rstrip() for line in text.splitlines())
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        pages.append({
            "n": i,
            "image": "slides/%s/slide-%02d.png" % (slug, i),
            "width": page.rect.width * ZOOM,
            "height": page.rect.height * ZOOM,
            "pdf_text": text,
            "bytes": os.path.getsize(png),
        })
    n = len(pages)
    doc.close()
    return pages, n


def extract_pptx(path):
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = None
        bullets = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            txt = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
            if not txt:
                continue
            if shape == slide.shapes.title and title is None:
                title = txt
            else:
                bullets.append(txt)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:  # noqa: BLE001
            notes = ""
        slides.append({"n": idx, "title": title, "bullets": bullets, "notes": notes})
    return slides


def main():
    decks = []
    pdfs = sorted(glob.glob(os.path.join(DECKS, "*.pdf")))
    for pdf in pdfs:
        name = os.path.basename(pdf)
        stem = name[:-4]
        slug = deck_slug(stem)
        t0 = time.time()
        pages, n = render_pdf(pdf, slug)
        pptx_path = os.path.join(DECKS, stem + ".pptx")
        pptx = []
        if os.path.exists(pptx_path):
            try:
                pptx = extract_pptx(pptx_path)
            except Exception as exc:  # noqa: BLE001
                print("  pptx parse failed %s: %s" % (stem, exc), file=sys.stderr)
        # merge pptx text onto pages
        for p in pages:
            src = next((s for s in pptx if s["n"] == p["n"]), None)
            p["title"] = (src or {}).get("title")
            p["bullets"] = (src or {}).get("bullets", [])
            p["notes"] = (src or {}).get("notes", "")
        total = sum(p["bytes"] for p in pages)
        decks.append({
            "stem": stem,
            "slug": slug,
            "pdf": "slides/downloads/%s" % name,
            "pptx": "slides/downloads/%s.pptx" % stem if os.path.exists(pptx_path) else None,
            "pages": pages,
            "page_count": len(pages),
            "pptx_slide_count": len(pptx),
            "bytes": total,
        })
        print("  %-34s %3d pages  %6.1f MB  %4.1fs" % (slug, len(pages), total / 1e6, time.time() - t0))

    with open(os.path.join(DATA, "slides.json"), "w") as fh:
        json.dump(decks, fh, indent=1)
    print("\ndecks %d  pages %d  total %.1f MB" % (
        len(decks), sum(d["page_count"] for d in decks),
        sum(d["bytes"] for d in decks) / 1e6))


if __name__ == "__main__":
    main()
